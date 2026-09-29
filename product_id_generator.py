import os
import sys
import sqlite3
import jdatetime
import openpyxl
import tkinter as tk
from tkinter import ttk, messagebox, filedialog, simpledialog

DEFAULT_PRODUCTS = [
    ("110101", "سبک سازمانی"),
    ("210201", "مانده بدهکار سازمانی"),
    ("310301", "عمر و حوادث گروهی")
]
RESET_PASSWORD = "1174"

def get_today_jalali():
    return jdatetime.date.today().strftime("%Y/%m/%d")


def get_data_dir():
    base = os.environ.get("LOCALAPPDATA", os.path.expanduser("~"))
    p = os.path.join(base, "PersianProductID")
    os.makedirs(p, exist_ok=True)
    return p


class Database:
    def __init__(self):
        self.db_path = os.path.join(get_data_dir(), "records.sqlite3")
        self.init_db()

    def get_conn(self):
        conn = sqlite3.connect(self.db_path)
        conn.row_factory = sqlite3.Row
        return conn

    def init_db(self):
        with self.get_conn() as conn:
            conn.execute("""
                CREATE TABLE IF NOT EXISTS settings (
                    k TEXT PRIMARY KEY,
                    v TEXT
                )
            """)
            conn.execute("""
                CREATE TABLE IF NOT EXISTS products (
                    code TEXT PRIMARY KEY,
                    name TEXT,
                    deleted INTEGER DEFAULT 0
                )
            """)
            conn.execute("""
                CREATE TABLE IF NOT EXISTS records (
                    id INTEGER PRIMARY KEY AUTOINCREMENT,
                    year INTEGER,
                    pc TEXT,
                    pname TEXT,
                    seq INTEGER,
                    full TEXT UNIQUE,
                    descr TEXT,
                    created TEXT,
                    deleted INTEGER DEFAULT 0,
                    UNIQUE(year, pc, seq)
                )
            """)
            try:
                conn.execute("ALTER TABLE products ADD COLUMN deleted INTEGER DEFAULT 0")
            except Exception:
                pass

            conn.execute("INSERT OR IGNORE INTO settings VALUES (?, ?)", ("start_year", "1400"))
            for code, name in DEFAULT_PRODUCTS:
                conn.execute("INSERT OR IGNORE INTO products (code, name, deleted) VALUES (?, ?, 0)", (code, name))
            conn.commit()

    def get_start_year(self):
        with self.get_conn() as conn:
            row = conn.execute("SELECT v FROM settings WHERE k=?", ("start_year",)).fetchone()
            return int(row["v"]) if row else 1400

    def get_products(self):
        with self.get_conn() as conn:
            return conn.execute("SELECT code, name FROM products WHERE deleted=0 ORDER BY code").fetchall()

    def save_settings(self, start_year, products_list):
        if not (start_year.isdigit() and len(start_year) == 4):
            raise ValueError("سال شروع باید یک عدد ۴ رقمی معتبر باشد.")

        seen = set()
        for code, name in products_list:
            if not (1 <= len(code) <= 10):
                raise ValueError(f"شناسه محصول '{code}' باید بین ۱ تا ۱۰ کاراکتر باشد.")
            if not name.strip():
                raise ValueError("نام محصول نمی‌تواند خالی باشد.")
            if code in seen:
                raise ValueError(f"شناسه محصول تکراری است: {code}")
            seen.add(code)

        with self.get_conn() as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                conn.execute("UPDATE products SET deleted=1")
                for code, name in products_list:
                    conn.execute("""
                        INSERT INTO products (code, name, deleted) VALUES (?, ?, 0)
                        ON CONFLICT(code) DO UPDATE SET name=excluded.name, deleted=0
                    """, (code, name.strip()))
                conn.execute("REPLACE INTO settings VALUES (?, ?)", ("start_year", start_year))
                conn.commit()
            except Exception:
                conn.rollback()
                raise

    def propose_next_id(self, product_code):
        year = self.get_start_year()
        with self.get_conn() as conn:
            row = conn.execute("SELECT MAX(seq) AS m FROM records WHERE year=? AND pc=?", (year, product_code)).fetchone()
            max_seq = row["m"] or 0

        if max_seq >= 9999:
            raise ValueError("ظرفیت شماره‌گذاری (۹۹۹۹) برای این محصول در سال جاری تکمیل شده است.")

        next_seq = max_seq + 1
        full_id = f"{year}/{product_code}/{next_seq:04d}"
        return full_id, year, next_seq

    def add_record(self, year, product_code, product_name, seq, descr):
        if not descr.strip():
            raise ValueError("توضیحات الزامی است.")

        with self.get_conn() as conn:
            prod = conn.execute("SELECT 1 FROM products WHERE code=? AND deleted=0", (product_code,)).fetchone()
            if not prod:
                raise ValueError("محصول انتخاب‌شده معتبر یا فعال نیست.")

            if year < self.get_start_year():
                raise ValueError("سال نمی‌تواند قبل از سال شروع سیستم باشد.")

            existing = conn.execute("SELECT 1 FROM records WHERE year=? AND pc=? AND seq=?", (year, product_code, seq)).fetchone()
            if existing:
                raise ValueError("این رکورد قبلاً در سیستم ثبت شده است.")

            full_id = f"{year}/{product_code}/{seq:04d}"
            conn.execute("""
                INSERT INTO records (year, pc, pname, seq, full, descr, created, deleted)
                VALUES (?, ?, ?, ?, ?, ?, ?, 0)
            """, (year, product_code, product_name, seq, full_id, descr.strip(), get_today_jalali()))
            conn.commit()

    def get_records(self, archived=False, f_code=None, f_desc=None, d1=None, d2=None, pcs=None):
        query = ["SELECT * FROM records WHERE deleted=?"]
        params = [1 if archived else 0]

        if f_code:
            query.append("AND full LIKE ?")
            params.append(f"%{f_code}%")
        if f_desc:
            query.append("AND descr LIKE ?")
            params.append(f"%{f_desc}%")
        if d1:
            query.append("AND created >= ?")
            params.append(d1)
        if d2:
            query.append("AND created <= ?")
            params.append(d2)
        if pcs is not None:
            if pcs:
                placeholders = ",".join("?" for _ in pcs)
                query.append(f"AND pc IN ({placeholders})")
                params.extend(pcs)
            else:
                query.append("AND 0")

        query.append("ORDER BY id DESC")
        with self.get_conn() as conn:
            return conn.execute(" ".join(query), params).fetchall()

    def get_counts(self):
        with self.get_conn() as conn:
            active_cnt = conn.execute("SELECT COUNT(*) AS c FROM records WHERE deleted=0").fetchone()["c"]
            arch_cnt = conn.execute("SELECT COUNT(*) AS c FROM records WHERE deleted=1").fetchone()["c"]
            return active_cnt, arch_cnt

    def set_archived(self, record_ids, archived=True):
        if not record_ids:
            return
        val = 1 if archived else 0
        placeholders = ",".join("?" for _ in record_ids)
        with self.get_conn() as conn:
            conn.execute(f"UPDATE records SET deleted=? WHERE id IN ({placeholders})", [val] + list(record_ids))
            conn.commit()

    def reset_main(self):
        with self.get_conn() as conn:
            conn.execute("DELETE FROM records WHERE deleted=0")
            conn.commit()

    def reset_archive(self):
        with self.get_conn() as conn:
            conn.execute("DELETE FROM records WHERE deleted=1")
            conn.commit()

    def reset_all(self):
        with self.get_conn() as conn:
            conn.execute("BEGIN IMMEDIATE")
            try:
                conn.execute("DROP TABLE IF EXISTS records")
                conn.execute("DROP TABLE IF EXISTS products")
                conn.execute("DROP TABLE IF EXISTS settings")
                conn.commit()
            except Exception:
                conn.rollback()
                raise
        self.init_db()


class AskDescription(tk.Toplevel):
    def __init__(self, parent, full_id):
        super().__init__(parent)
        self.title("ورود توضیحات شناسه")
        self.geometry("520x300")
        self.resizable(False, False)
        self.transient(parent)
        self.grab_set()

        self.result = None

        ttk.Label(self, text=f"شناسه پیشنهادی: {full_id}", font=("Consolas", 12, "bold"), foreground="#004d40").pack(pady=(12, 4))
        ttk.Label(self, text="لطفاً توضیحات کامل را وارد نمایید (الزامی):", font=("Segoe UI", 9, "bold")).pack(anchor="e", padx=15, pady=(4, 2))

        self.txt = tk.Text(self, wrap="word", height=8, font=("Segoe UI", 10), relief="solid", bd=1)
        self.txt.pack(fill="both", expand=True, padx=15, pady=6)
        self.txt.focus_set()

        btn_box = ttk.Frame(self)
        btn_box.pack(fill="x", padx=15, pady=(0, 12))
        ttk.Button(btn_box, text="تأیید و ثبت", command=self.on_ok).pack(side="left", padx=4)
        ttk.Button(btn_box, text="انصراف", command=self.destroy).pack(side="left", padx=4)

        self.wait_window(self)

    def on_ok(self):
        val = self.txt.get("1.0", "end").strip()
        if not val:
            messagebox.showwarning("خطا", "توضیحات نمی‌تواند خالی باشد.", parent=self)
            return
        self.result = val
        self.destroy()


class FilterPanel(ttk.Frame):
    def __init__(self, master, on_change, get_products_cb):
        super().__init__(master)
        self.on_change = on_change
        self.get_products_cb = get_products_cb

        f = ttk.Frame(self)
        f.pack(fill="x", padx=6, pady=4)

        ttk.Label(f, text="کد:").pack(side="right", padx=2)
        self.e_code = ttk.Entry(f, width=14, justify="right")
        self.e_code.pack(side="right", padx=3)

        ttk.Label(f, text="توضیحات:").pack(side="right", padx=2)
        self.e_desc = ttk.Entry(f, width=18, justify="right")
        self.e_desc.pack(side="right", padx=3)

        ttk.Label(f, text="از تاریخ:").pack(side="right", padx=2)
        self.e_d1 = ttk.Entry(f, width=10, justify="right")
        self.e_d1.pack(side="right", padx=3)

        ttk.Label(f, text="تا تاریخ:").pack(side="right", padx=2)
        self.e_d2 = ttk.Entry(f, width=10, justify="right")
        self.e_d2.pack(side="right", padx=3)

        self.mb = ttk.Menubutton(f, text="محصولات: همه")
        self.mb.pack(side="right", padx=4)

        ttk.Button(f, text="جستجو", command=self.on_change).pack(side="right", padx=3)
        ttk.Button(f, text="حذف فیلتر", command=self.clear).pack(side="right", padx=3)

        self.vars = {}
        self.rebuild_menu()

    def rebuild_menu(self):
        # Keep each code's prior state when rebuilding (e.g. after refresh).
        previous = {code: var.get() for code, var in self.vars.items()}
        products = self.get_products_cb()
        menu = tk.Menu(self.mb, tearoff=0)
        self.mb["menu"] = menu
        self.vars = {}
        for r in products:
            c, n = r["code"], r["name"]
            v = tk.BooleanVar(value=previous.get(c, True))
            self.vars[c] = v
            menu.add_checkbutton(label=f"{c} - {n}", variable=v, command=lambda c=c: self.toggle(c))
        self.all_var = tk.BooleanVar(value=all(v.get() for v in self.vars.values()))
        menu.insert_checkbutton(0, label="همه", variable=self.all_var, command=self.select_all)
        self.update_caption()

    def select_all(self):
        v = self.all_var.get()
        for var in self.vars.values():
            var.set(v)
        self.update_caption()
        self.on_change()

    def toggle(self, _):
        all_checked = all(var.get() for var in self.vars.values())
        self.all_var.set(all_checked)
        self.update_caption()
        self.on_change()

    def update_caption(self):
        sel = [c for c, v in self.vars.items() if v.get()]
        if len(sel) == len(self.vars) or not self.vars:
            self.mb.config(text="محصولات: همه")
        elif not sel:
            self.mb.config(text="محصولات: هیچ‌کدام")
        else:
            self.mb.config(text=f"محصولات: ({len(sel)}) انتخاب شده")

    def values(self):
        sel = [c for c, v in self.vars.items() if v.get()] if self.vars else None
        return {
            "f_code": self.e_code.get().strip() or None,
            "f_desc": self.e_desc.get().strip() or None,
            "d1": self.e_d1.get().strip() or None,
            "d2": self.e_d2.get().strip() or None,
            "pcs": sel
        }

    def clear(self):
        self.e_code.delete(0, "end")
        self.e_desc.delete(0, "end")
        self.e_d1.delete(0, "end")
        self.e_d2.delete(0, "end")
        self.all_var.set(True)
        self.select_all()


def build_table(parent):
    cols = ("chk", "full", "pname", "descr", "created", "status")
    tree = ttk.Treeview(parent, columns=cols, show="headings", selectmode="extended")
    tree.heading("chk", text="[✓]")
    tree.heading("full", text="شناسه")
    tree.heading("pname", text="محصول")
    tree.heading("descr", text="توضیحات")
    tree.heading("created", text="تاریخ")
    tree.heading("status", text="وضعیت")

    tree.column("chk", width=45, anchor="center", stretch=False)
    tree.column("full", width=160, anchor="center")
    tree.column("pname", width=170, anchor="center")
    tree.column("descr", width=340, anchor="e")
    tree.column("created", width=110, anchor="center")
    tree.column("status", width=90, anchor="center")

    sb = ttk.Scrollbar(parent, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=sb.set)
    tree.pack(side="left", fill="both", expand=True)
    sb.pack(side="right", fill="y")

    def on_click(event):
        item = tree.identify_row(event.y)
        col = tree.identify_column(event.x)
        if item and col == "#1":
            cur = tree.item(item, "values")
            if cur:
                new_state = "☐" if cur[0] == "☑" else "☑"
                tree.item(item, values=(new_state, *cur[1:]))

    tree.bind("<Button-1>", on_click)
    tree.bind("<Double-1>", lambda e: copy_selected_ids(tree))
    return tree


def get_checked_ids(tree):
    checked = []
    for item in tree.get_children():
        vals = tree.item(item, "values")
        if vals and vals[0] == "☑":
            checked.append(item)
    if checked:
        return checked
    return list(tree.selection())


def copy_selected_ids(tree):
    target = get_checked_ids(tree)
    ids = [str(tree.item(i, "values")[1]) for i in target if tree.item(i, "values")]
    if not ids:
        messagebox.showinfo("کپی", "هیچ رکوردی انتخاب نشده است.")
        return
    tree.clipboard_clear()
    tree.clipboard_append("\n".join(ids))
    messagebox.showinfo("کپی", f"{len(ids)} شناسه در کلیپ‌بورد کپی شد.")


def export_excel(parent, rows):
    if not rows:
        messagebox.showinfo("خروجی اکسل", "داده‌ای برای خروجی وجود ندارد.")
        return
    path = filedialog.asksaveasfilename(parent=parent, defaultextension=".xlsx", filetypes=[("Excel Files", "*.xlsx")])
    if not path:
        return
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "شناسه‌ها"
    ws.views.sheetView[0].rightToLeft = True
    ws.append(["شناسه", "محصول", "توضیحات", "تاریخ", "وضعیت"])
    for r in rows:
        st = "بایگانی" if r["deleted"] else "فعال"
        ws.append([r["full"], r["pname"], r["descr"], r["created"], st])
    wb.save(path)
    messagebox.showinfo("خروجی اکسل", f"{len(rows)} رکورد با موفقیت ذخیره شد.")

def get_resource_path(relative_path):
    """ دریافت مسیر فایل برای حالت exe و حالت عادی """
    if hasattr(sys, '_MEIPASS'):
        # این مسیر جایی است که PyInstaller فایل‌ها را در حالت onefile اکسترکت می‌کند
        return os.path.join(sys._MEIPASS, relative_path)
    return os.path.join(os.path.abspath("."), relative_path)

icon_path = get_resource_path("logo.ico")



class ArchiveWindow(tk.Toplevel):
    def __init__(self, master, db):
        super().__init__(master)
        self.db = db
        self.title("بایگانی شناسه‌ها")
        self.geometry("1060x640")
        self.iconbitmap(icon_path)
        self.transient(master)

        self.fp = FilterPanel(self, self.refresh, self.db.get_products)
        self.fp.pack(fill="x")

        table_frame = ttk.Frame(self)
        table_frame.pack(fill="both", expand=True, padx=8, pady=4)
        self.tree = build_table(table_frame)

        bf = ttk.Frame(self)
        bf.pack(fill="x", padx=8, pady=6)

        ttk.Button(bf, text="انتخاب همه", command=lambda: self.set_all_checks(True)).pack(side="right", padx=3)
        ttk.Button(bf, text="لغو انتخاب همه", command=lambda: self.set_all_checks(False)).pack(side="right", padx=3)
        ttk.Button(bf, text="بازگردانی به لیست اصلی", command=self.restore_selected).pack(side="right", padx=5)
        ttk.Button(bf, text="کپی شناسه‌های انتخاب‌شده", command=lambda: copy_selected_ids(self.tree)).pack(side="right", padx=3)
        ttk.Button(bf, text="خروجی اکسل", command=self.export).pack(side="right", padx=3)

        self.lbl_count = ttk.Label(bf, text="تعداد کل رکوردهای بایگانی‌شده: ۰", font=("Segoe UI", 9, "bold"))
        self.lbl_count.pack(side="left", padx=8)

        self.refresh()

    def set_all_checks(self, check):
        val = "☑" if check else "☐"
        for i in self.tree.get_children():
            cur = self.tree.item(i, "values")
            if cur:
                self.tree.item(i, values=(val, *cur[1:]))

    def refresh(self):
        self.current_rows = self.db.get_records(archived=True, **self.fp.values())
        self.tree.delete(*self.tree.get_children())
        for r in self.current_rows:
            self.tree.insert("", "end", iid=str(r["id"]), values=("☐", r["full"], r["pname"], r["descr"], r["created"], "بایگانی"))
        _, arch_cnt = self.db.get_counts()
        self.lbl_count.config(text=f"تعداد کل رکوردهای بایگانی‌شده: {arch_cnt}")

    def restore_selected(self):
        target = get_checked_ids(self.tree)
        if not target:
            messagebox.showinfo("بازگردانی", "هیچ رکوردی برای بازگردانی انتخاب نشده است.", parent=self)
            return
        if not messagebox.askyesno("تأیید بازگردانی", f"آیا از بازگردانی {len(target)} رکورد به لیست اصلی اطمینان دارید؟", parent=self):
            return
        self.db.set_archived(target, archived=False)
        self.refresh()
        self.master.refresh()
        messagebox.showinfo("موفقیت", "رکوردهای انتخاب‌شده با موفقیت به لیست اصلی بازگردانده شدند.", parent=self)

    def export(self):
        export_excel(self, self.current_rows)


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.db = Database()
        self.title("نرم افزار تولید و مدیریت شناسه")
        self.geometry("1060x670")
        self.iconbitmap(icon_path)

        self.build_menu()
        self.build_ui()
        self.refresh()

    def build_menu(self):
        mb = tk.Menu(self)
        self.config(menu=mb)


        m_set = tk.Menu(mb, tearoff=0)
        mb.add_cascade(label="تنظیمات", menu=m_set)
        m_set.add_command(label="تنظیمات سیستم", command=self.open_settings)
        m_set.add_command(label="ریست جدول اصلی", command=self.reset_main)
        m_set.add_command(label="ریست بایگانی", command=self.reset_archive)
        m_set.add_separator()
        m_set.add_command(label="ریست کلی دیتابیس", command=self.reset_all)

        m_tools = tk.Menu(mb, tearoff=0)
        mb.add_cascade(label="ابزارها", menu=m_tools)
        m_tools.add_command(label="بایگانی", command=self.open_archive)
        
        m_about = tk.Menu(mb, tearoff=0)
        mb.add_cascade(label="درباره", menu=m_about)
        m_about.add_command(label="درباره توسعه‌دهنده", command=self.show_about)

    def build_ui(self):
        top = ttk.Frame(self)
        top.pack(fill="x", padx=10, pady=8)

        self.cb_prod = ttk.Combobox(top, state="readonly", width=32)
        self.cb_prod.pack(side="right", padx=5)

        ttk.Button(top, text="تولید شناسه جدید", command=self.generate_id).pack(side="right", padx=5)

        # دکمه کپی شناسه در کنار لیبل شناسه پیشنهادی
        self.btn_copy_prop = ttk.Button(top, text="کپی شناسه", command=self.copy_proposed_id)
        self.btn_copy_prop.pack(side="right", padx=4)

        self.lbl_prop = ttk.Label(top, text="----/------/----", font=("Consolas", 14, "bold"), foreground="#004d40")
        self.lbl_prop.pack(side="right", padx=6)

        self.fp = FilterPanel(self, self.refresh, self.db.get_products)
        self.fp.pack(fill="x", padx=4)

        table_frame = ttk.Frame(self)
        table_frame.pack(fill="both", expand=True, padx=10, pady=4)
        self.tree = build_table(table_frame)

        bf = ttk.Frame(self)
        bf.pack(fill="x", padx=10, pady=6)

        ttk.Button(bf, text="انتخاب همه", command=lambda: self.set_all_checks(True)).pack(side="right", padx=3)
        ttk.Button(bf, text="لغو انتخاب همه", command=lambda: self.set_all_checks(False)).pack(side="right", padx=3)
        ttk.Button(bf, text="کپی شناسه‌های انتخاب‌شده", command=lambda: copy_selected_ids(self.tree)).pack(side="right", padx=3)
        ttk.Button(bf, text="انتقال به بایگانی", command=self.archive_selected).pack(side="right", padx=3)
        ttk.Button(bf, text="خروجی اکسل", command=self.export).pack(side="right", padx=3)

        self.lbl_counts = ttk.Label(bf, text="تعداد رکوردهای اصلی: ۰ | بایگانی‌شده: ۰", font=("Segoe UI", 9, "bold"))
        self.lbl_counts.pack(side="left", padx=8)

    def set_all_checks(self, check):
        val = "☑" if check else "☐"
        for i in self.tree.get_children():
            cur = self.tree.item(i, "values")
            if cur:
                self.tree.item(i, values=(val, *cur[1:]))

    def show_about(self):
        msg = "Author : Amin Ahankoubi\nEmail : Amin.ahankobi@gmail.com"
        messagebox.showinfo("درباره توسعه‌دهنده", msg, parent=self)

    def open_archive(self):
        ArchiveWindow(self, self.db)

    def refresh(self):
        prods = self.db.get_products()
        self.cb_prod["values"] = [f"{r['code']} - {r['name']}" for r in prods]
        if prods and not self.cb_prod.get():
            self.cb_prod.current(0)
        self.fp.rebuild_menu()

        self.current_rows = self.db.get_records(archived=False, **self.fp.values())
        self.tree.delete(*self.tree.get_children())
        for r in self.current_rows:
            self.tree.insert("", "end", iid=str(r["id"]), values=("☐", r["full"], r["pname"], r["descr"], r["created"], "فعال"))

        active_cnt, arch_cnt = self.db.get_counts()
        self.lbl_counts.config(text=f"تعداد رکوردهای اصلی: {active_cnt} | بایگانی‌شده: {arch_cnt}")

    def copy_proposed_id(self):
        val = self.lbl_prop.cget("text")
        if not val or "-" in val:
            messagebox.showinfo("کپی", "هنوز شناسه‌ای تولید نشده است.")
            return
        self.clipboard_clear()
        self.clipboard_append(val)
        messagebox.showinfo("کپی", f"شناسه {val} کپی شد.")

    def generate_id(self):
        sel = self.cb_prod.get()
        if not sel:
            messagebox.showwarning("خطا", "لطفاً ابتدا محصول را مشخص نمایید.")
            return
        code = sel.split(" - ")[0]
        pname = sel.split(" - ")[1]

        try:
            prop_id, year, seq = self.db.propose_next_id(code)
        except Exception as e:
            messagebox.showerror("خطا", str(e))
            return

        self.lbl_prop.config(text=prop_id)

        dlg = AskDescription(self, prop_id)
        desc = dlg.result
        if not desc:
            return

        if messagebox.askyesno("تأیید نهایی ثبت", f"شناسه تولیدشده: {prop_id}\n\nآیا از ثبت نهایی این شناسه اطمینان دارید؟"):
            try:
                self.db.add_record(year, code, pname, seq, desc)
                self.refresh()
                messagebox.showinfo("موفقیت", f"شناسه {prop_id} با موفقیت در سیستم ثبت گردید.")
            except Exception as e:
                messagebox.showerror("خطا در ثبت", str(e))

    def archive_selected(self):
        target = get_checked_ids(self.tree)
        if not target:
            messagebox.showinfo("انتقال به بایگانی", "هیچ رکوردی انتخاب نشده است.")
            return
        if not messagebox.askyesno("تأیید انتقال", f"آیا از انتقال {len(target)} رکورد به بایگانی مطمئن هستید؟"):
            return
        self.db.set_archived(target, archived=True)
        self.refresh()
        messagebox.showinfo("موفقیت", "رکوردهای انتخاب‌شده به بایگانی منتقل شدند.")

    def export(self):
        export_excel(self, self.current_rows)

    def reset_main(self):
        if messagebox.askyesno("ریست جدول اصلی", "آیا از پاک‌سازی تمامی رکوردهای جدول اصلی اطمینان دارید؟\nاین عملیات غیرقابل بازگشت است."):
            self.db.reset_main()
            self.refresh()
            messagebox.showinfo("موفقیت", "رکوردهای جدول اصلی به طور کامل حذف شدند.")

    def reset_archive(self):
        if messagebox.askyesno("ریست بایگانی", "آیا از پاک‌سازی تمامی رکوردهای بخش بایگانی اطمینان دارید؟\nاین عملیات غیرقابل بازگشت است."):
            self.db.reset_archive()
            self.refresh()
            messagebox.showinfo("موفقیت", "رکوردهای بخش بایگانی به طور کامل حذف شدند.")

    def reset_all(self):
        if not messagebox.askyesno("هشدار بسیار مهم", "آیا از ریست کلی مطمئن هستید؟\nتمامی داده‌ها، جداول و تنظیمات به طور کامل حذف خواهند شد!"):
            return

        password = simpledialog.askstring("رمز ریست کلی", "رمز چهار رقمی را وارد کنید:", show="*", parent=self)
        if password is None:
            return
        if password != RESET_PASSWORD:
            messagebox.showerror("خطا", "رمز نادرست است.", parent=self)
            return

        try:
            self.db.reset_all()
            self.refresh()
            messagebox.showinfo("موفقیت", "پایگاه‌داده به طور کامل بازنشانی و ریست شد.")
        except Exception as e:
            messagebox.showerror("خطا", f"خطا در ریست پایگاه‌داده: {e}")

    def open_settings(self):
        win = tk.Toplevel(self)
        win.title("تنظیمات سیستم")
        win.geometry("560x520")
        win.iconbitmap(icon_path)
        win.transient(self)
        win.grab_set()

        top_f = ttk.Frame(win)
        top_f.pack(fill="x", padx=10, pady=8)
        ttk.Label(top_f, text="سال شروع سیستم:").pack(side="right", padx=5)
        e_year = ttk.Entry(top_f, width=8, justify="center")
        e_year.pack(side="right", padx=5)
        e_year.insert(0, str(self.db.get_start_year()))

        cols = ("code", "name")
        tree = ttk.Treeview(win, columns=cols, show="headings", height=9)
        tree.heading("code", text="شناسه محصول (۱ تا ۱۰ رقم)")
        tree.heading("name", text="نام محصول")
        tree.column("code", width=160, anchor="center")
        tree.column("name", width=340, anchor="e")
        tree.pack(fill="both", expand=True, padx=10, pady=5)

        for r in self.db.get_products():
            tree.insert("", "end", values=(r["code"], r["name"]))

        inp = ttk.Frame(win)
        inp.pack(fill="x", padx=10, pady=4)
        ttk.Label(inp, text="شناسه:").pack(side="right", padx=2)
        e_c = ttk.Entry(inp, width=12, justify="center")
        e_c.pack(side="right", padx=4)
        ttk.Label(inp, text="نام:").pack(side="right", padx=2)
        e_n = ttk.Entry(inp, width=22, justify="right")
        e_n.pack(side="right", padx=4)

        def pick(_):
            sel = tree.selection()
            if sel:
                v = tree.item(sel[0], "values")
                e_c.delete(0, "end")
                e_c.insert(0, v[0])
                e_n.delete(0, "end")
                e_n.insert(0, v[1])

        tree.bind("<<TreeviewSelect>>", pick)

        def add_product():
            c = e_c.get().strip()
            n = e_n.get().strip()
            if not (1 <= len(c) <= 10):
                messagebox.showwarning("خطا", "شناسه محصول باید بین ۱ تا ۱۰ کاراکتر باشد.", parent=win)
                return
            if not n:
                messagebox.showwarning("خطا", "نام محصول الزامی است.", parent=win)
                return

            for item in tree.get_children():
                if tree.item(item, "values")[0] == c:
                    messagebox.showwarning("خطا", "این شناسه محصول از قبل در لیست وجود دارد. برای تغییر نام از دکمه ویرایش استفاده کنید.", parent=win)
                    return

            tree.insert("", "end", values=(c, n))
            e_c.delete(0, "end")
            e_n.delete(0, "end")

        def update_product():
            c = e_c.get().strip()
            n = e_n.get().strip()
            if not (1 <= len(c) <= 10):
                messagebox.showwarning("خطا", "شناسه محصول باید بین ۱ تا ۱۰ کاراکتر باشد.", parent=win)
                return
            if not n:
                messagebox.showwarning("خطا", "نام محصول الزامی است.", parent=win)
                return

            sel = tree.selection()
            target = None
            if sel:
                target = sel[0]
            else:
                for item in tree.get_children():
                    if tree.item(item, "values")[0] == c:
                        target = item
                        break
            if not target:
                messagebox.showwarning("خطا", "لطفاً ابتدا یک محصول را از جدول انتخاب کنید یا شناسه معتبر وارد نمایید.", parent=win)
                return

            for item in tree.get_children():
                if item != target and tree.item(item, "values")[0] == c:
                    messagebox.showwarning("خطا", "این شناسه محصول از قبل در لیست وجود دارد. برای تغییر نام از دکمه ویرایش استفاده کنید.", parent=win)
                    return

            tree.item(target, values=(c, n))
            e_c.delete(0, "end")
            e_n.delete(0, "end")

        def delete_item():
            sel = tree.selection()
            if not sel:
                return
            tree.delete(sel[0])

        btn_grid = ttk.Frame(win)
        btn_grid.pack(fill="x", padx=10, pady=4)
        ttk.Button(btn_grid, text="افزودن", command=add_product).pack(side="right", padx=4)
        ttk.Button(btn_grid, text="ویرایش محصول", command=update_product).pack(side="right", padx=4)
        ttk.Button(btn_grid, text="حذف محصول از لیست", command=delete_item).pack(side="right", padx=4)

        def save_all():
            prods = [tree.item(i, "values") for i in tree.get_children()]
            try:
                self.db.save_settings(e_year.get().strip(), prods)
                self.refresh()
                messagebox.showinfo("موفقیت", "تنظیمات با موفقیت ذخیره گردید.", parent=win)
                win.destroy()
            except Exception as ex:
                messagebox.showerror("خطا", str(ex), parent=win)

        bf2 = ttk.Frame(win)
        bf2.pack(fill="x", padx=10, pady=8)
        ttk.Button(bf2, text="ثبت نهایی تنظیمات", command=save_all).pack(side="right", padx=5)
        ttk.Button(bf2, text="خروج", command=win.destroy).pack(side="left", padx=5)


if __name__ == "__main__":
    App().mainloop()
